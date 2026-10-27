"""Выгрузка списка пользователей в Excel (только для администратора)."""
from __future__ import annotations

import io

from tests.conftest import auth_headers, make_user


async def test_admin_downloads_users_list(client, db, admin_user):
    await make_user(
        db, username="ivanov", email="ivanov@example.com", full_name="Иванов Иван Иванович",
    )
    r = await client.get("/admin/export/users", headers=auth_headers(admin_user))
    assert r.status_code == 200
    assert "aegis_users_" in r.headers["content-disposition"]

    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:3] == ("ФИО", "Логин", "Почта")
    flat = [c for row in rows for c in row if isinstance(c, str)]
    assert "Иванов Иван Иванович" in flat
    assert "ivanov" in flat
    assert "ivanov@example.com" in flat
    # Хеш пароля не должен утекать в файл.
    assert not any(c.startswith("$2") for c in flat)


async def test_formula_in_full_name_is_escaped(client, db, admin_user):
    await make_user(db, username="evil2", full_name="=1+1")
    r = await client.get("/admin/export/users", headers=auth_headers(admin_user))
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(r.content)).active
    names = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert "'=1+1" in names


async def test_reader_cannot_export_users(client, db):
    reader = await make_user(db, username="plainreader")
    r = await client.get("/admin/export/users", headers=auth_headers(reader))
    assert r.status_code == 403
